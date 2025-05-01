"""
Tool abstractions and example tool implementations with async support.
"""
from typing import List
import asyncio
import os
import aiohttp
from docx import Document
import tempfile
import datetime

class Tool:
    async def arun(self, query: str, memory: List[str]) -> str:
        raise NotImplementedError

class SearchTool(Tool):
    async def arun(self, query: str, memory: List[str]) -> str:
        api_key = os.getenv("SERPAPI_API_KEY")
        if not api_key:
            return "Search API key not set. Set SERPAPI_API_KEY environment variable."
        params = {
            "q": query,
            "api_key": api_key,
            "engine": "google"
        }
        url = "https://serpapi.com/search"
        async with aiohttp.ClientSession() as session:
            async with session.get(url, params=params) as resp:
                if resp.status != 200:
                    return f"Search API error: {resp.status}"
                data = await resp.json()
                # Get the first organic result
                results = data.get("organic_results", [])
                if results:
                    return results[0].get("snippet") or results[0].get("title") or str(results[0])
                return "No search results found."

class WeatherTool(Tool):
    async def arun(self, query: str, memory: List[str]) -> str:
        api_key = os.getenv("OPENWEATHER_API_KEY")
        if not api_key:
            return "Weather API key not set. Set OPENWEATHER_API_KEY environment variable."
        # Try to extract city from query (very basic)
        import re
        match = re.search(r'in ([A-Za-z ]+)', query)
        city = match.group(1).strip() if match else "London"
        url = f"https://api.openweathermap.org/data/2.5/weather"
        params = {"q": city, "appid": api_key, "units": "metric"}
        async with aiohttp.ClientSession() as session:
            async with session.get(url, params=params) as resp:
                if resp.status != 200:
                    return f"Weather API error: {resp.status}"
                data = await resp.json()
                weather = data.get("weather", [{}])[0].get("description", "No data")
                temp = data.get("main", {}).get("temp", "?")
                return f"Weather in {city}: {weather}, {temp}°C"

class DocumentationTool(Tool):
    def __init__(self, llm):
        self.llm = llm

    async def arun(self, query: str, memory: List[str]) -> str:
        prompt = f"Generate documentation or explanation for the following request/code/query:\n{query}\nPlease provide clear, concise, and well-formatted documentation."
        doc_text = await self.llm.agenerate(prompt, memory)
        try:
            # Create a Word document in the project 'generated_docs' folder
            project_dir = os.path.dirname(os.path.abspath(__file__))
            docs_dir = os.path.join(project_dir, "generated_docs")
            os.makedirs(docs_dir, exist_ok=True)
            timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            file_path = os.path.join(docs_dir, f"doc_{timestamp}.docx")
            doc = Document()
            doc.add_heading('Generated Documentation', 0)
            for para in doc_text.split('\n'):
                doc.add_paragraph(para)
            doc.save(file_path)
            return f"Documentation generated.\n\n{doc_text}\n\nWord document saved at: {file_path}"
        except Exception as e:
            return f"Documentation generated but failed to save Word document. Error: {e}\n\n{doc_text}"
